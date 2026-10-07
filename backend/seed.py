from datetime import date, datetime
import os
from app import create_app
from app.extensions import db
from app.models.user import User
from app.models.department import (
    DepartmentInfo,
    FacultyMember,
    DepartmentEvent,
    StudentAchievement,
    Announcement,
)

app = create_app("development")

with app.app_context():
    print("Initializing database tables...")
    db.create_all()

    # 1. Seed Administrative & HOD Users
    hod_email = "hod@department.edu"
    admin_email = "admin@department.edu"

    if not User.query.filter_by(email=hod_email).first():
        hod = User(
            email=hod_email,
            full_name="Dr. M. A. Jabbar",
            role="HOD",
            department="Computer Science & Engineering",
            is_active=True,
        )
        hod.set_password("Admin@123")
        db.session.add(hod)
        print("Created HOD User: hod@department.edu / Admin@123")

    if not User.query.filter_by(email=admin_email).first():
        admin = User(
            email=admin_email,
            full_name="System Administrator",
            role="ADMIN",
            department="Computer Science & Engineering",
            is_active=True,
        )
        admin.set_password("Admin@123")
        db.session.add(admin)
        print("Created Admin User: admin@department.edu / Admin@123")

    # 2. Seed Department Info
    if not DepartmentInfo.query.first():
        dept = DepartmentInfo(
            name="Department of Computer Science & Engineering",
            short_code="CSE",
            tagline="Fostering Academic Rigor, Advanced Research & Technological Innovation",
            about=(
                "The Department of Computer Science & Engineering is committed to delivering "
                "world-class technical education. With state-of-the-art computational infrastructure, "
                "distinguished faculty, and an outcome-based educational paradigm, the department "
                "equips students with the analytical capability, systems thinking, and ethical foundation "
                "required to lead in computing technologies and academic research."
            ),
            vision=(
                "To be a premier center of academic excellence and pioneering research in Computer "
                "Science and Engineering, producing globally competent professionals who innovate with social responsibility."
            ),
            mission=(
                "1. Provide a rigorous, contemporary curriculum grounded in foundational science and modern computing practices.\n"
                "2. Cultivate research, critical inquiry, and interdisciplinary collaboration across faculty and student bodies.\n"
                "3. Foster ethical leadership, continuous self-learning, and societal technological engagement."
            ),
            hod_name="Dr. M. A. Jabbar",
            hod_designation="Professor & Head of Department",
            hod_qualification="Ph.D. (CSE), M.Tech, B.Tech, SMIEEE, FIETE",
            hod_message=(
                "Welcome to the Department of Computer Science & Engineering. Our primary objective is to build "
                "an intellectually vibrant environment that bridges rigorous academic theory with real-world technical execution. "
                "Through continuous assessment, subject-wise attendance analytics, and outcome-oriented mentorship, we monitor "
                "every student's trajectory to ensure no learner is left behind while top performers are propelled toward research "
                "and industry leadership. I invite our students and faculty to uphold the highest benchmarks of discipline and academic integrity."
            ),
            hod_photo_url="https://images.unsplash.com/photo-1534528741775-53994a69daeb?auto=format&fit=crop&w=400&q=80",
            contact_email="hod.cse@college.edu",
            contact_phone="+91 40 2345 6789",
            office_location="Block-A, Third Floor, Academic Enclave",
            established_year=2008,
        )
        db.session.add(dept)
        print("Created Department Info.")

    # 3. Seed Faculty Members
    if FacultyMember.query.count() == 0:
        faculty_data = [
            {
                "name": "Dr. M. A. Jabbar",
                "designation": "Professor & HOD",
                "qualification": "Ph.D. (CSE), M.Tech",
                "specialization": "Machine Learning, Network Security & Big Data",
                "email": "hod.cse@college.edu",
                "phone": "+91 40 2345 6789",
                "experience_years": 22,
                "display_order": 1,
                "bio": "Senior Member IEEE and Fellow IETE with over 22 years in academic administration, artificial intelligence, and network vulnerability research.",
            },
            {
                "name": "Dr. K. Radhika",
                "designation": "Professor",
                "qualification": "Ph.D. (CSE), M.Tech",
                "specialization": "Data Mining, Knowledge Discovery & AI Ethics",
                "email": "k.radhika@college.edu",
                "phone": "+91 40 2345 6790",
                "experience_years": 18,
                "display_order": 2,
                "bio": "Specialist in large-scale data systems and predictive models with 45+ international journal citations.",
            },
            {
                "name": "Dr. S. Suresh",
                "designation": "Associate Professor",
                "qualification": "Ph.D. (CSE)",
                "specialization": "Distributed Systems, Cloud Architecture & Microservices",
                "email": "s.suresh@college.edu",
                "phone": "+91 40 2345 6791",
                "experience_years": 14,
                "display_order": 3,
                "bio": "Focuses on high-availability backend architectures, fault-tolerant consensus mechanisms, and containerized virtualization.",
            },
            {
                "name": "Dr. P. Anitha",
                "designation": "Associate Professor",
                "qualification": "Ph.D. (CSE)",
                "specialization": "Natural Language Processing & Deep Learning",
                "email": "p.anitha@college.edu",
                "phone": "+91 40 2345 6792",
                "experience_years": 12,
                "display_order": 4,
                "bio": "Active researcher in speech translation architectures and low-resource multilingual LLM fine-tuning.",
            },
            {
                "name": "Mr. V. Rajesh",
                "designation": "Assistant Professor",
                "qualification": "M.Tech (Software Engineering)",
                "specialization": "Computer Vision & Autonomous Systems",
                "email": "v.rajesh@college.edu",
                "phone": "+91 40 2345 6793",
                "experience_years": 8,
                "display_order": 5,
                "bio": "Leads the departmental Autonomous Robotics & Embedded Perception Lab.",
            },
            {
                "name": "Mrs. T. Sunitha",
                "designation": "Assistant Professor",
                "qualification": "M.Tech (CSE)",
                "specialization": "IoT Sensors, Edge Computing & Embedded OS",
                "email": "t.sunitha@college.edu",
                "phone": "+91 40 2345 6794",
                "experience_years": 7,
                "display_order": 6,
                "bio": "Oversees hardware prototyping, smart agricultural telemetry, and real-time edge streaming pipelines.",
            },
            {
                "name": "Mr. B. Mahesh",
                "designation": "Assistant Professor",
                "qualification": "M.Tech (Cybersecurity)",
                "specialization": "Cryptographic Protocols & Reverse Engineering",
                "email": "b.mahesh@college.edu",
                "phone": "+91 40 2345 6795",
                "experience_years": 6,
                "display_order": 7,
                "bio": "Advises the college Capture-The-Flag (CTF) security team and teaches operating systems internals.",
            },
            {
                "name": "Ms. N. Kavitha",
                "designation": "Assistant Professor",
                "qualification": "M.Tech (Computer Science)",
                "specialization": "Database Management Systems & Relational Theory",
                "email": "n.kavitha@college.edu",
                "phone": "+91 40 2345 6796",
                "experience_years": 5,
                "display_order": 8,
                "bio": "Leads 2nd Year academic coordination, lab curriculum synchronization, and database performance tuning sessions.",
            },
        ]
        for f in faculty_data:
            db.session.add(FacultyMember(**f))
        print("Seeded 8 Faculty Members.")

    # 4. Seed Department Events
    if DepartmentEvent.query.count() == 0:
        events_data = [
            {
                "title": "National Conference on Emerging AI Frontiers (NC-EAIF)",
                "description": "Peer-reviewed national symposium bringing together academic researchers and industry practitioners to address scalable neural architectures and safe computing.",
                "category": "Conference",
                "event_date": date(2026, 11, 14),
                "location": "Auditorium Complex & Seminar Hall 1",
                "is_featured": True,
                "is_upcoming": True,
            },
            {
                "title": "36-Hour Hackathon: Smart Campus & Edge IoT",
                "description": "Intensive team hackathon challenging undergraduate students to design intelligent campus energy, security, and academic monitoring systems.",
                "category": "Hackathon",
                "event_date": date(2026, 10, 28),
                "location": "Advanced Computing Laboratory (Block A)",
                "is_featured": True,
                "is_upcoming": True,
            },
            {
                "title": "Faculty Development Program: Cloud Infrastructure & Kubernetes",
                "description": "One-week hands-on FDP on container orchestration, microservice telemetry, and continuous deployment workflows.",
                "category": "Workshop",
                "event_date": date(2026, 9, 18),
                "location": "Virtual Seminar Room",
                "is_featured": False,
                "is_upcoming": False,
            },
            {
                "title": "Industry Keynote: High-Performance Database Systems",
                "description": "Guest lecture delivered by Principal Architect at AWS discussing distributed transactional isolation and storage engine optimizations.",
                "category": "Guest Lecture",
                "event_date": date(2026, 8, 22),
                "location": "Main CSE Seminar Hall",
                "is_featured": False,
                "is_upcoming": False,
            },
        ]
        for e in events_data:
            db.session.add(DepartmentEvent(**e))
        print("Seeded Department Events.")

    # 5. Seed Student Achievements
    if StudentAchievement.query.count() == 0:
        achievements_data = [
            {
                "title": "Smart India Hackathon 2025 - 1st Runner Up",
                "student_names": "Aditya Sharma, Sneha Rao, Rohan Verma, Tanvi Patel",
                "category": "Hackathon",
                "event_name": "Smart India Hackathon (Ministry of Education)",
                "award": "1st Runner Up (Rs. 75,000 Cash Prize)",
                "achievement_date": date(2025, 12, 19),
                "description": "Developed an offline-first mobile telemetry system for rural health clinic supply tracking.",
            },
            {
                "title": "ACM ICPC Regional Finalists Rank #14",
                "student_names": "K. Sai Praneeth, R. Karthik, M. Ananya",
                "category": "Competitive Programming",
                "event_name": "ACM ICPC Amritapuri Regional",
                "award": "Regional Rank 14",
                "achievement_date": date(2025, 11, 10),
                "description": "Solved 8 complex algorithmic problems within 5 hours against top collegiate teams nationwide.",
            },
            {
                "title": "IEEE Best Undergraduate Research Paper",
                "student_names": "Divya Krishnan, Rahul Nambiar",
                "category": "Research Publication",
                "event_name": "IEEE International Conference on Data Science (ICDS)",
                "award": "Best Undergraduate Research Paper",
                "achievement_date": date(2025, 10, 5),
                "description": "Published research on lightweight quantization techniques for embedded vision models.",
            },
        ]
        for a in achievements_data:
            db.session.add(StudentAchievement(**a))
        print("Seeded Student Achievements.")

    # 6. Seed Announcements
    if Announcement.query.count() == 0:
        announcements_data = [
            {
                "title": "End-Semester Academic & Lab Assessment Timetable",
                "content": "The schedule for 2nd Year Semesters III & IV end-term practical examinations and internal theory assessments has been officially released. All students must ensure their laboratory observation logs are certified prior to the examinations.",
                "category": "Examination",
                "is_pinned": True,
                "publish_date": date(2026, 10, 5),
            },
            {
                "title": "Mandatory Subject-Wise Attendance Review for 2nd Year Sections A, B & C",
                "content": "Notice to all 2nd Year students: Institutional regulation mandates a minimum of 75% subject-wise attendance to qualify for semester examinations. Students in the condonation band (65%-75%) must report to their academic section coordinators.",
                "category": "Academic",
                "is_pinned": True,
                "publish_date": date(2026, 10, 2),
            },
            {
                "title": "Commencement of Doubt-Clearing & Remedial Modules",
                "content": "Special remedial classes for Mathematics-III, Data Structures, and Computer Organization will begin this Saturday. Check the department notice board for section timings.",
                "category": "Notice",
            },
        ]
        for ann in announcements_data:
            db.session.add(Announcement(**ann))
        print("Seeded Announcements.")

    # 7. Seed Phase 2 Academic Hierarchy (Batch 2025-2029)
    from app.models.academic import (
        Batch,
        AcademicYear,
        Semester,
        Section,
        Student,
        Subject,
    )

    batch_name = "2025-2029"
    batch = Batch.query.filter_by(name=batch_name).first()
    if not batch:
        batch = Batch(
            name=batch_name,
            start_year=2025,
            end_year=2029,
            is_active=True,
        )
        db.session.add(batch)
        db.session.flush()
        print(f"Created Batch: {batch_name}")

        year_names = ["1st Year", "2nd Year", "3rd Year", "4th Year"]
        academic_years = []
        for idx, y_name in enumerate(year_names, start=1):
            cal_start = 2025 + (idx - 1)
            ay = AcademicYear(
                batch_id=batch.id,
                year_number=idx,
                name=y_name,
                calendar_year=f"{cal_start}-{cal_start + 1}",
                is_current=(idx == 2),  # 2nd Year is current focus
            )
            db.session.add(ay)
            academic_years.append(ay)
        db.session.flush()
        print("Created 4 Academic Years.")

        semesters = []
        for ay in academic_years:
            s1_num = (ay.year_number * 2) - 1
            s2_num = ay.year_number * 2

            s1 = Semester(
                academic_year_id=ay.id,
                semester_number=s1_num,
                name=f"Semester {s1_num}",
                is_current=(s1_num == 3),  # Semester 3 is current focus
            )
            s2 = Semester(
                academic_year_id=ay.id,
                semester_number=s2_num,
                name=f"Semester {s2_num}",
                is_current=False,
            )
            db.session.add_all([s1, s2])
            semesters.extend([s1, s2])
        db.session.flush()
        print("Created 8 Semesters.")

        # Create Sections for Semesters 1, 2, 3
        section_map = {}
        for s in semesters:
            if s.semester_number in [1, 2, 3]:
                for sec_letter in ["A", "B", "C"]:
                    sec = Section(
                        semester_id=s.id,
                        name=sec_letter,
                        room_number=f"Room-{sec_letter}-30{s.semester_number}",
                    )
                    db.session.add(sec)
                    section_map[(s.semester_number, sec_letter)] = sec
        db.session.flush()
        print("Created Configurable Sections (A, B, C) for Semesters 1, 2, 3.")

        # Seed Subjects for Semester 3
        sem3 = next((s for s in semesters if s.semester_number == 3), None)
        if sem3:
            sem3_subjects = [
                {"code": "A9002", "name": "Ordinary Differential Equations and Calculus of Variations", "short_name": "ODECV", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9009", "name": "Engineering Chemistry", "short_name": "EC", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9011", "name": "Engineering Science Elective", "short_name": "ESE", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9503", "name": "Data Structures using C++", "short_name": "DS", "credits": 4.0, "subject_type": "THEORY"},
                {"code": "A9402", "name": "Digital Electronics", "short_name": "DE", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9010", "name": "Engineering Chemistry Lab", "short_name": "ECL", "credits": 1.5, "subject_type": "LAB"},
                {"code": "A9012", "name": "Engineering Science Elective Lab", "short_name": "ESEL", "credits": 1.5, "subject_type": "LAB"},
                {"code": "A9504", "name": "Data Structures Lab", "short_name": "DSL", "credits": 1.5, "subject_type": "LAB"},
                {"code": "A9304", "name": "Computer Aided Engineering Graphics", "short_name": "CAEG", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9022", "name": "Professional Development & Design", "short_name": "PDD", "credits": 2.0, "subject_type": "THEORY"},
            ]
            for sub in sem3_subjects:
                db.session.add(Subject(semester_id=sem3.id, **sub))
            print(f"Created {len(sem3_subjects)} Subjects for Semester 3.")

        # Seed Subjects for Semester 1 (for 1-1 results)
        sem1 = next((s for s in semesters if s.semester_number == 1), None)
        if sem1:
            sem1_subjects = [
                {"code": "A9001", "name": "Matrices and Calculus", "short_name": "MAC", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9501", "name": "Programming for Problem Solving", "short_name": "PPS", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9502", "name": "Programming for Problem Solving Lab", "short_name": "PPSL", "credits": 1.5, "subject_type": "LAB"},
                {"code": "A9302", "name": "Engineering Workshop", "short_name": "EW", "credits": 2.5, "subject_type": "LAB"},
                {"code": "A9021", "name": "Critical Thinking & Design Thinking", "short_name": "CCDT", "credits": 2.0, "subject_type": "THEORY"},
                {"code": "A9007", "name": "Engineering Physics", "short_name": "EP", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9008", "name": "Engineering Physics Lab", "short_name": "EPL", "credits": 1.5, "subject_type": "LAB"},
                {"code": "A9204", "name": "Basic Electrical Engineering", "short_name": "BEE", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9801", "name": "Foundations of Data Science", "short_name": "FDS", "credits": 3.0, "subject_type": "THEORY"},
                {"code": "A9205", "name": "Basic Electrical Engineering Lab", "short_name": "BEEL", "credits": 1.5, "subject_type": "LAB"},
            ]
            for sub in sem1_subjects:
                db.session.add(Subject(semester_id=sem1.id, **sub))
            print(f"Created {len(sem1_subjects)} Subjects for Semester 1.")

        # Seed Students from Department Roster
        import openpyxl
        roster_files = [
            ("A", os.path.join("..", "Students data", "ATTENDENCE", "7. CSM-I B.Tech. II _A.xlsx")),
            ("B", os.path.join("..", "Students data", "ATTENDENCE", "7. CSM-I B.Tech. II _B.xlsx")),
            ("C", os.path.join("..", "Students data", "ATTENDENCE", "7. CSM-I B.Tech. II _C.xlsx")),
        ]

        total_students_seeded = 0
        for sec_letter, rel_path in roster_files:
            target_sec = section_map.get((3, sec_letter))
            if os.path.exists(rel_path):
                try:
                    wb = openpyxl.load_workbook(rel_path, data_only=True)
                    ws = wb.active
                    for r in range(7, ws.max_row + 1):
                        roll_val = ws.cell(r, 3).value
                        if roll_val and str(roll_val).strip() and not any(w in str(roll_val).lower() for w in ["total", "avg", "percent"]):
                            roll = str(roll_val).strip().upper()
                            name_val = ws.cell(r, 4).value
                            name = str(name_val).strip() if name_val else f"Student {roll}"
                            
                            st = Student(
                                roll_number=roll,
                                name=name,
                                batch_id=batch.id,
                                current_section_id=target_sec.id if target_sec else None,
                                email=f"{roll.lower()}@department.edu",
                                is_active=True,
                            )
                            db.session.add(st)
                            total_students_seeded += 1
                except Exception as e:
                    print(f"Notice: Could not parse {rel_path}: {e}")

        # Fallback if roster files were not located
        if total_students_seeded == 0:
            for i in range(1, 66):
                roll = f"25881A66{i:02d}"
                st = Student(
                    roll_number=roll,
                    name=f"Student {roll}",
                    batch_id=batch.id,
                    current_section_id=section_map.get((3, "A")).id if section_map.get((3, "A")) else None,
                    is_active=True,
                )
                db.session.add(st)
                total_students_seeded += 1

        print(f"Seeded {total_students_seeded} Students for Batch 2025-2029 across Sections A, B, C.")

    db.session.commit()
    print("Database seeding completed successfully!")

