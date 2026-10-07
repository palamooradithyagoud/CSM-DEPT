"""
Excel (.xlsx) Report Generator for Academic Management.
Generates styled, institutional spreadsheets with openpyxl for Department, Section, Student, Subject, and Insights reports.
"""

import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelReportGenerator:
    """Generates styled institutional Excel workbooks using openpyxl."""

    HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")  # Dark Slate
    SECTION_FILL = PatternFill(start_color="1E3A2F", end_color="1E3A2F", fill_type="solid") # Deep Forest/Academic Emerald
    SUBHEADER_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    STRIPE_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    ALERT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="0F172A")
    FONT_REGULAR = Font(name="Calibri", size=11, color="1E293B")
    FONT_MUTED = Font(name="Calibri", size=10, italic=True, color="64748B")

    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    @classmethod
    def _auto_fit_columns(cls, ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ws.merged_cells:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def _create_workbook(cls):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.views.sheetView[0].showGridLines = True
        return wb, ws

    @classmethod
    def _to_bytes(cls, wb):
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    @classmethod
    def generate_department_excel(cls, data):
        wb, ws = cls._create_workbook()
        ws.title = "Department Overview"

        # 1. Main Header Title
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING — ACADEMIC REPORT"
        title_cell.font = cls.FONT_TITLE
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = cls.ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        # Meta info
        meta_rows = [
            ("Batch", data.get("batchName", "All Batches"), "Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")),
            ("Semester", data.get("semesterName", "All Semesters"), "Status", "Official Institutional Record"),
        ]
        r = 2
        for row in meta_rows:
            ws.cell(row=r, column=1, value=row[0]).font = cls.FONT_BOLD
            ws.cell(row=r, column=2, value=row[1]).font = cls.FONT_REGULAR
            ws.cell(row=r, column=4, value=row[2]).font = cls.FONT_BOLD
            ws.cell(row=r, column=5, value=row[3]).font = cls.FONT_REGULAR
            r += 1

        r += 1
        # 2. Executive KPIs Table
        ws.merge_cells(f"A{r}:F{r}")
        sec_cell = ws[f"A{r}"]
        sec_cell.value = "EXECUTIVE ACADEMIC PERFORMANCE METRICS"
        sec_cell.font = cls.FONT_HEADER
        sec_cell.fill = cls.SECTION_FILL
        sec_cell.alignment = cls.ALIGN_LEFT
        ws.row_dimensions[r].height = 24
        r += 1

        kpis = data.get("kpis", {})
        kpi_items = [
            ("Total Enrolled Students", kpis.get("totalStudents", 0)),
            ("Average SGPA", kpis.get("averageSGPA", "N/A")),
            ("Average CGPA", kpis.get("averageCGPA", "N/A")),
            ("Average Attendance", f"{kpis.get('averageAttendance', 'N/A')}%" if kpis.get('averageAttendance') is not None else "N/A"),
            ("Pass Percentage", f"{kpis.get('passPercentage', 'N/A')}%" if kpis.get('passPercentage') is not None else "N/A"),
            ("Students Requiring Attention", data.get("flaggedStudentsCount", 0)),
        ]
        ws.cell(row=r, column=1, value="Metric").font = cls.FONT_BOLD
        ws.cell(row=r, column=2, value="Value").font = cls.FONT_BOLD
        ws.cell(row=r, column=1).border = cls.THIN_BORDER
        ws.cell(row=r, column=2).border = cls.THIN_BORDER
        r += 1

        for label, val in kpi_items:
            c1 = ws.cell(row=r, column=1, value=label)
            c2 = ws.cell(row=r, column=2, value=val)
            c1.font = cls.FONT_REGULAR
            c2.font = cls.FONT_BOLD
            c1.border = cls.THIN_BORDER
            c2.border = cls.THIN_BORDER
            if label == "Students Requiring Attention" and val > 0:
                c2.fill = cls.ALERT_FILL
            r += 1

        r += 1
        # 3. Subject Performance Breakdown
        ws.merge_cells(f"A{r}:F{r}")
        ws[f"A{r}"].value = "SUBJECT-WISE PERFORMANCE BREAKDOWN"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        ws.row_dimensions[r].height = 24
        r += 1

        headers = ["Subject Code", "Subject Name", "Avg Marks", "Pass %", "Failures", "Total Evaluated"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.alignment = cls.ALIGN_CENTER
            c.border = cls.THIN_BORDER
        r += 1

        for sub in data.get("subjects", []):
            vals = [
                sub.get("code", ""),
                sub.get("name", ""),
                sub.get("averageMarks", "N/A"),
                f"{sub.get('passPercentage', 'N/A')}%" if sub.get('passPercentage') is not None else "N/A",
                sub.get("failureCount", 0),
                sub.get("evaluatedStudents", 0)
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
                c.alignment = cls.ALIGN_CENTER if col_idx in [1, 3, 4, 5, 6] else cls.ALIGN_LEFT
            r += 1

        r += 1
        # 4. Section Comparison
        ws.merge_cells(f"A{r}:F{r}")
        ws[f"A{r}"].value = "SECTION COMPARATIVE METRICS"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        ws.row_dimensions[r].height = 24
        r += 1

        sec_headers = ["Section", "Student Count", "Average SGPA", "Average Attendance (%)", "Pass Rate (%)"]
        for col_idx, h in enumerate(sec_headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.alignment = cls.ALIGN_CENTER
            c.border = cls.THIN_BORDER
        r += 1

        for sec in data.get("sections", []):
            vals = [
                sec.get("sectionName", ""),
                sec.get("studentCount", 0),
                sec.get("averageSGPA", "N/A"),
                f"{sec.get('averageAttendance', 'N/A')}%" if sec.get('averageAttendance') is not None else "N/A",
                f"{sec.get('passPercentage', 'N/A')}%" if sec.get('passPercentage') is not None else "N/A"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
                c.alignment = cls.ALIGN_CENTER
            r += 1

        cls._auto_fit_columns(ws)
        return cls._to_bytes(wb)

    @classmethod
    def generate_section_excel(cls, data):
        wb, ws = cls._create_workbook()
        ws.title = f"Section {data.get('sectionName', '')}"

        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = f"SECTION {data.get('sectionName', '')} — ACADEMIC PERFORMANCE REPORT"
        title_cell.font = cls.FONT_TITLE
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = cls.ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        # Meta
        ws.cell(row=2, column=1, value="Batch:").font = cls.FONT_BOLD
        ws.cell(row=2, column=2, value=data.get("batchName", "")).font = cls.FONT_REGULAR
        ws.cell(row=2, column=4, value="Semester:").font = cls.FONT_BOLD
        ws.cell(row=2, column=5, value=data.get("semesterName", "")).font = cls.FONT_REGULAR

        r = 4
        # Section KPIs
        ws.merge_cells(f"A{r}:F{r}")
        ws[f"A{r}"].value = "SECTION SUMMARY METRICS"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        sec_kpis = [
            ("Enrolled Students", data.get("studentCount", 0)),
            ("Average SGPA", data.get("averageSGPA", "N/A")),
            ("Average Attendance", f"{data.get('averageAttendance', 'N/A')}%" if data.get('averageAttendance') is not None else "N/A"),
            ("Pass Rate", f"{data.get('passPercentage', 'N/A')}%" if data.get('passPercentage') is not None else "N/A"),
            ("Students Requiring Attention", data.get("flaggedStudentsCount", 0)),
        ]
        for label, val in sec_kpis:
            c1 = ws.cell(row=r, column=1, value=label)
            c2 = ws.cell(row=r, column=2, value=val)
            c1.font = cls.FONT_REGULAR
            c2.font = cls.FONT_BOLD
            c1.border = cls.THIN_BORDER
            c2.border = cls.THIN_BORDER
            r += 1

        r += 1
        # Subject Breakdown
        ws.merge_cells(f"A{r}:F{r}")
        ws[f"A{r}"].value = "SUBJECT PERFORMANCE IN SECTION"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        headers = ["Code", "Subject Name", "Avg Marks", "Pass %", "Failures", "Total Evaluated"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.border = cls.THIN_BORDER
        r += 1

        for sub in data.get("subjects", []):
            vals = [
                sub.get("code", ""),
                sub.get("name", ""),
                sub.get("averageMarks", "N/A"),
                f"{sub.get('passPercentage', 'N/A')}%" if sub.get('passPercentage') is not None else "N/A",
                sub.get("failureCount", 0),
                sub.get("evaluatedStudents", 0)
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
            r += 1

        r += 1
        # Students Requiring Attention
        if data.get("flaggedStudents"):
            ws.merge_cells(f"A{r}:F{r}")
            ws[f"A{r}"].value = "STUDENTS REQUIRING ACADEMIC ATTENTION"
            ws[f"A{r}"].font = cls.FONT_HEADER
            ws[f"A{r}"].fill = cls.SECTION_FILL
            r += 1

            att_headers = ["Roll Number", "Student Name", "Severity", "Primary Reason", "Attendance", "SGPA"]
            for col_idx, h in enumerate(att_headers, 1):
                c = ws.cell(row=r, column=col_idx, value=h)
                c.font = cls.FONT_HEADER
                c.fill = cls.SUBHEADER_FILL
                c.border = cls.THIN_BORDER
            r += 1

            for st in data.get("flaggedStudents", []):
                vals = [
                    st.get("rollNumber", ""),
                    st.get("name", ""),
                    st.get("severity", ""),
                    st.get("primaryReason", ""),
                    f"{st.get('attendance', 'N/A')}%" if st.get('attendance') is not None else "N/A",
                    st.get("sgpa", "N/A")
                ]
                for col_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=r, column=col_idx, value=val)
                    c.font = cls.FONT_REGULAR
                    c.border = cls.THIN_BORDER
                    if col_idx == 3 and val in ["CRITICAL", "HIGH"]:
                        c.fill = cls.ALERT_FILL
                r += 1

        cls._auto_fit_columns(ws)
        return cls._to_bytes(wb)

    @classmethod
    def generate_student_excel(cls, data):
        wb, ws = cls._create_workbook()
        ws.title = f"Student {data.get('rollNumber', '')}"

        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"STUDENT ACADEMIC DOSSIER — {data.get('name', '').upper()}"
        title_cell.font = cls.FONT_TITLE
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = cls.ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        # Meta info
        meta = [
            ("Roll Number", data.get("rollNumber", ""), "Batch", data.get("batchName", "")),
            ("Current Section", data.get("sectionName", ""), "Semester", data.get("semesterName", "")),
            ("Latest SGPA", data.get("sgpa", "N/A"), "Overall CGPA", data.get("cgpa", "N/A")),
            ("Average Attendance", f"{data.get('attendancePercentage', 'N/A')}%" if data.get('attendancePercentage') is not None else "N/A", "Email", data.get("email", "N/A")),
        ]
        r = 2
        for row in meta:
            ws.cell(row=r, column=1, value=row[0]).font = cls.FONT_BOLD
            ws.cell(row=r, column=2, value=row[1]).font = cls.FONT_REGULAR
            ws.cell(row=r, column=4, value=row[2]).font = cls.FONT_BOLD
            ws.cell(row=r, column=5, value=row[3]).font = cls.FONT_REGULAR
            r += 1

        r += 1
        # Subject-wise table
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"].value = "COURSEWORK RESULTS & ATTENDANCE"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        headers = ["Code", "Subject Name", "Credits", "Attendance", "Total Marks", "Grade", "Status"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.border = cls.THIN_BORDER
        r += 1

        for sub in data.get("subjects", []):
            vals = [
                sub.get("code", ""),
                sub.get("name", ""),
                sub.get("credits", ""),
                f"{sub.get('attendance', 'N/A')}%" if sub.get('attendance') is not None else "N/A",
                sub.get("totalMarks", "N/A"),
                sub.get("grade", "N/A"),
                sub.get("resultStatus", "N/A")
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
                if col_idx == 7 and val == "FAILED":
                    c.fill = cls.ALERT_FILL
            r += 1

        r += 1
        # Academic Trajectory
        if data.get("trajectory"):
            ws.merge_cells(f"A{r}:D{r}")
            ws[f"A{r}"].value = "SEMESTER-WISE PROGRESSION"
            ws[f"A{r}"].font = cls.FONT_HEADER
            ws[f"A{r}"].fill = cls.SECTION_FILL
            r += 1

            traj_headers = ["Semester", "SGPA", "CGPA", "Credits Earned"]
            for col_idx, h in enumerate(traj_headers, 1):
                c = ws.cell(row=r, column=col_idx, value=h)
                c.font = cls.FONT_HEADER
                c.fill = cls.SUBHEADER_FILL
                c.border = cls.THIN_BORDER
            r += 1

            for sem in data.get("trajectory", []):
                vals = [sem.get("semesterName", ""), sem.get("sgpa", "N/A"), sem.get("cgpa", "N/A"), sem.get("credits", "N/A")]
                for col_idx, val in enumerate(vals, 1):
                    c = ws.cell(row=r, column=col_idx, value=val)
                    c.font = cls.FONT_REGULAR
                    c.border = cls.THIN_BORDER
                r += 1

        cls._auto_fit_columns(ws)
        return cls._to_bytes(wb)

    @classmethod
    def generate_subject_excel(cls, data):
        wb, ws = cls._create_workbook()
        ws.title = f"Subject {data.get('code', '')}"

        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = f"SUBJECT PERFORMANCE REPORT — {data.get('code', '')} {data.get('name', '')}"
        title_cell.font = cls.FONT_TITLE
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = cls.ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        # Meta
        ws.cell(row=2, column=1, value="Credits:").font = cls.FONT_BOLD
        ws.cell(row=2, column=2, value=data.get("credits", "")).font = cls.FONT_REGULAR
        ws.cell(row=2, column=4, value="Subject Type:").font = cls.FONT_BOLD
        ws.cell(row=2, column=5, value=data.get("subjectType", "")).font = cls.FONT_REGULAR

        r = 4
        ws.merge_cells(f"A{r}:F{r}")
        ws[f"A{r}"].value = "OVERALL METRICS"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        sub_kpis = [
            ("Enrolled / Evaluated Students", data.get("studentCount", 0)),
            ("Average Marks", data.get("averageMarks", "N/A")),
            ("Highest Marks", data.get("highestMarks", "N/A")),
            ("Lowest Marks", data.get("lowestMarks", "N/A")),
            ("Pass Rate", f"{data.get('passPercentage', 'N/A')}%" if data.get('passPercentage') is not None else "N/A"),
            ("Failures", data.get("failureCount", 0)),
            ("Average Attendance", f"{data.get('averageAttendance', 'N/A')}%" if data.get('averageAttendance') is not None else "N/A"),
        ]
        for label, val in sub_kpis:
            c1 = ws.cell(row=r, column=1, value=label)
            c2 = ws.cell(row=r, column=2, value=val)
            c1.font = cls.FONT_REGULAR
            c2.font = cls.FONT_BOLD
            c1.border = cls.THIN_BORDER
            c2.border = cls.THIN_BORDER
            r += 1

        r += 1
        # Section breakdown
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"].value = "SECTION-WISE COMPARISON"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        headers = ["Section", "Student Count", "Average Marks", "Pass Rate (%)", "Failures"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.border = cls.THIN_BORDER
        r += 1

        for sec in data.get("sections", []):
            vals = [
                sec.get("sectionName", ""),
                sec.get("studentCount", 0),
                sec.get("averageMarks", "N/A"),
                f"{sec.get('passPercentage', 'N/A')}%" if sec.get('passPercentage') is not None else "N/A",
                sec.get("failureCount", 0)
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
            r += 1

        cls._auto_fit_columns(ws)
        return cls._to_bytes(wb)

    @classmethod
    def generate_insights_excel(cls, data):
        wb, ws = cls._create_workbook()
        ws.title = "Academic Watchlist"

        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "ACADEMIC INSIGHTS & STUDENT WATCHLIST"
        title_cell.font = cls.FONT_TITLE
        title_cell.fill = cls.HEADER_FILL
        title_cell.alignment = cls.ALIGN_CENTER
        ws.row_dimensions[1].height = 36

        r = 3
        # Summary
        summary = data.get("summary", {})
        ws.merge_cells(f"A{r}:D{r}")
        ws[f"A{r}"].value = "PROBLEM IDENTIFICATION SUMMARY"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        items = [
            ("Students Requiring Attention", summary.get("studentsRequiringAttention", 0)),
            ("Critical Priority Issues", summary.get("criticalIssues", 0)),
            ("High Priority Issues", summary.get("highPriorityIssues", 0)),
            ("Subjects Flagged", summary.get("subjectsRequiringAttention", 0)),
            ("Sections Flagged", summary.get("sectionsRequiringAttention", 0)),
        ]
        for label, val in items:
            c1 = ws.cell(row=r, column=1, value=label)
            c2 = ws.cell(row=r, column=2, value=val)
            c1.font = cls.FONT_REGULAR
            c2.font = cls.FONT_BOLD
            c1.border = cls.THIN_BORDER
            c2.border = cls.THIN_BORDER
            r += 1

        r += 1
        # Watchlist
        ws.merge_cells(f"A{r}:G{r}")
        ws[f"A{r}"].value = "STUDENT DIAGNOSTIC WATCHLIST"
        ws[f"A{r}"].font = cls.FONT_HEADER
        ws[f"A{r}"].fill = cls.SECTION_FILL
        r += 1

        headers = ["Roll Number", "Student Name", "Section", "Severity", "Primary Reason", "Attendance", "SGPA"]
        for col_idx, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=col_idx, value=h)
            c.font = cls.FONT_HEADER
            c.fill = cls.SUBHEADER_FILL
            c.border = cls.THIN_BORDER
        r += 1

        for st in data.get("watchlist", []):
            vals = [
                st.get("rollNumber", ""),
                st.get("name", ""),
                st.get("sectionName", ""),
                st.get("severity", ""),
                st.get("primaryReason", ""),
                f"{st.get('attendance', 'N/A')}%" if st.get('attendance') is not None else "N/A",
                st.get("sgpa", "N/A")
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws.cell(row=r, column=col_idx, value=val)
                c.font = cls.FONT_REGULAR
                c.border = cls.THIN_BORDER
                if col_idx == 4 and val in ["CRITICAL", "HIGH"]:
                    c.fill = cls.ALERT_FILL
            r += 1

        cls._auto_fit_columns(ws)
        return cls._to_bytes(wb)
